"""National registers four more EU regulators publish whole.

    Netherlands  CBG's Geneesmiddeleninformatiebank as one pipe-separated file,
                 weekly; holder, RVG/RVH or EU number, procedure, ATC, and the
                 SmPC and leaflet each product page links to.
    Poland       URPL's Register of Medicinal Products on the Ministry of
                 Health's register platform, daily (CC BY 4.0 on dane.gov.pl);
                 holder, permit number and validity, manufacturer and country.
    Czechia      SUKL's DLP database, monthly (SUKL open data); product,
                 composition, holder and country, with SUKL's own English
                 names for substances, forms, routes and countries.
    Denmark      DKMA's list of authorised medicines, daily; holder, ATC,
                 procedure. It carries no authorisation number.

Each is indexed locally like Italy's register (open_registers). What a
register publishes in its own language stays in that language unless the
register itself gives the English, or the term is a standard one listed
below; a term not listed keeps the regulator's wording rather than a guess.

Matching. Poland writes substances in pharmacopoeial Latin ("Metformini
hydrochloridum"), the Netherlands and Denmark glue the salt to the molecule
("METFORMINEHYDROCHLORIDE", "Atorvastatincalcium"). ``national_match_text``
undoes both on the register side only, so a search for "metformin" finds them
and the matching of every other register is untouched.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Iterator

from sources.open_registers import (
    OpenRegister,
    _fold,
    _match_text,
    add_register,
    download_file,
    search_register,
)
from sources.parser import extract_strength


def _clean(value: object) -> str:
    return " ".join(str(value or "").split())


def _unique(values: Iterable[str], joiner: str = "; ") -> str:
    return joiner.join(dict.fromkeys(value for value in (_clean(item) for item in values) if value))


# --------------------------------------------------------------- matching

# Salts, hydrates and counter-ions as Dutch, Danish and Latin write them when
# they are glued to the molecule's name. Longest first, so "dihydrochlorid"
# is taken before "hydrochlorid".
_GLUED_SUFFIXES = sorted(
    """
    hydrochloride hydrochlorid dihydrochloride dihydrochlorid hydrobromide hydrobromid
    besilaat besilat besylat maleaat maleat fumaraat fumarat tartraat tartrat
    mesilaat mesilat succinaat succinat acetaat acetat citraat citrat sulfaat sulfat
    fosfaat fosfat dinatrium natrium dikalium kalium calcium magnesium
    trihydraat trihydrat dihydraat dihydrat monohydraat monohydrat hemihydraat hemihydrat
    sesquihydraat sesquihydrat chloride chlorid bromide bromid zink zinc
    """.split(),
    key=len,
    reverse=True,
)
_MIN_STEM = 4
_VOWELS = set("aeiou")


def _unglue(word: str) -> list[str]:
    """ "metforminehydrochloride" -> ["metformine", "hydrochloride"]."""
    tail: list[str] = []
    changed = True
    while changed:
        changed = False
        for suffix in _GLUED_SUFFIXES:
            if word.endswith(suffix) and len(word) - len(suffix) >= _MIN_STEM:
                tail.insert(0, suffix)
                word = word[: -len(suffix)]
                changed = True
                break
    return [word, *tail]


# An acid glued to its name: Dutch "zoledroninezuur", Danish "zoledronsyre",
# both "zoledronic acid". The -ine/-in before it is the Dutch and Danish
# adjective ending, as in "valproïnezuur" and "fusidinsyre".
_GLUED_ACID = re.compile(r"^(?P<stem>[a-z]{4,}?)(?:ine|in)?(?:zuur|syre)$")


def _unglue_acid(word: str) -> list[str]:
    match = _GLUED_ACID.match(word)
    return [f"{match['stem']}ic", "acid"] if match else [word]


def _from_latin(word: str) -> str:
    """Pharmacopoeial Latin to the English stem: metformini, atorvastatinum -> metformin, atorvastatin."""
    if len(word) <= 5:
        return word
    if word.endswith("um"):
        return word[:-2]
    if word.endswith("i") and word[-2] not in _VOWELS:
        return word[:-1]
    return word


def national_match_text(text: object, latin: bool = False) -> str:
    words: list[str] = []
    for word in _fold(text).split():
        if latin:
            word = _from_latin(word)
        for part in _unglue_acid(word):
            words.extend(_unglue(part))
    return _match_text(" ".join(words))


def _indexed(match: str, row: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    """The row with the tokens it was matched on, folded the way the search
    page's relevance check compares them (as Russia's register does), so a
    Latin or glued name that matched the register also passes on the page."""
    from sources.grls_russia import fold

    row["inn_fold_tokens"] = " ".join(fold(token) for token in match.split())
    return match, row


def _date(value: object) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _clean(value)
    for pattern in ("%Y/%m/%d", "%Y-%m-%d", "%d.%m.%Y", "%d-%m-%Y", "%d/%m/%Y", "%y%m%d"):
        try:
            return datetime.strptime(text, pattern).date().isoformat()
        except ValueError:
            continue
    return text


# EDQM standard terms, for the forms that make up most of each register. A
# form not listed keeps the regulator's own wording.
DUTCH_FORMS = {
    "filmomhulde tablet": "Film-coated tablet",
    "tablet": "Tablet",
    "tablet met verlengde afgifte": "Prolonged-release tablet",
    "tablet met gereguleerde afgifte": "Modified-release tablet",
    "maagsapresistente tablet": "Gastro-resistant tablet",
    "orodispergeerbare tablet": "Orodispersible tablet",
    "kauwtablet": "Chewable tablet",
    "bruistablet": "Effervescent tablet",
    "dispergeerbare tablet": "Dispersible tablet",
    "capsule, hard": "Capsule, hard",
    "capsule, zacht": "Capsule, soft",
    "oplossing voor injectie": "Solution for injection",
    "oplossing voor infusie": "Solution for infusion",
    "concentraat voor oplossing voor infusie": "Concentrate for solution for infusion",
    "poeder en oplosmiddel voor oplossing voor injectie": "Powder and solvent for solution for injection",
    "poeder voor oplossing voor injectie": "Powder for solution for injection",
    "poeder voor concentraat voor oplossing voor infusie": "Powder for concentrate for solution for infusion",
    "suspensie voor injectie": "Suspension for injection",
    "drank": "Oral solution",
    "oogdruppels, oplossing": "Eye drops, solution",
    "neusspray, oplossing": "Nasal spray, solution",
    "zalf": "Ointment",
    "crème": "Cream",
    "gel": "Gel",
    "pleister voor transdermaal gebruik": "Transdermal patch",
}
POLISH_FORMS = {
    "tabletki powlekane": "Film-coated tablet",
    "tabletki": "Tablet",
    "tabletki o przedłużonym uwalnianiu": "Prolonged-release tablet",
    "tabletki o zmodyfikowanym uwalnianiu": "Modified-release tablet",
    "tabletki dojelitowe": "Gastro-resistant tablet",
    "tabletki ulegające rozpadowi w jamie ustnej": "Orodispersible tablet",
    "tabletki do rozgryzania i żucia": "Chewable tablet",
    "tabletki musujące": "Effervescent tablet",
    "kapsułki twarde": "Capsule, hard",
    "kapsułki miękkie": "Capsule, soft",
    "roztwór do wstrzykiwań": "Solution for injection",
    "roztwór do wstrzykiwań w ampułko-strzykawce": "Solution for injection in pre-filled syringe",
    "roztwór do infuzji": "Solution for infusion",
    "koncentrat do sporządzania roztworu do infuzji": "Concentrate for solution for infusion",
    "proszek i rozpuszczalnik do sporządzania roztworu do wstrzykiwań": "Powder and solvent for solution for injection",
    "proszek do sporządzania roztworu do wstrzykiwań": "Powder for solution for injection",
    "krople do oczu, roztwór": "Eye drops, solution",
    "roztwór doustny": "Oral solution",
    "zawiesina doustna": "Oral suspension",
    "krem": "Cream",
    "maść": "Ointment",
    "żel": "Gel",
    "system transdermalny": "Transdermal patch",
}
DANISH_FORMS = {
    "filmovertrukne tabletter": "Film-coated tablet",
    "tabletter": "Tablet",
    "depottabletter": "Prolonged-release tablet",
    "enterotabletter": "Gastro-resistant tablet",
    "smeltetabletter": "Orodispersible tablet",
    "tyggetabletter": "Chewable tablet",
    "brusetabletter": "Effervescent tablet",
    "kapsler, hårde": "Capsule, hard",
    "kapsler, bløde": "Capsule, soft",
    "depotkapsler, hårde": "Prolonged-release capsule, hard",
    "injektionsvæske, opløsning": "Solution for injection",
    "injektionsvæske, opløsning i fyldt injektionssprøjte": "Solution for injection in pre-filled syringe",
    "injektionsvæske, opløsning i fyldt pen": "Solution for injection in pre-filled pen",
    "injektionsvæske, suspension": "Suspension for injection",
    "infusionsvæske, opløsning": "Solution for infusion",
    "koncentrat til infusionsvæske, opløsning": "Concentrate for solution for infusion",
    "pulver og solvens til injektionsvæske, opløsning": "Powder and solvent for solution for injection",
    "pulver til injektionsvæske, opløsning": "Powder for solution for injection",
    "øjendråber, opløsning": "Eye drops, solution",
    "oral opløsning": "Oral solution",
    "oral suspension": "Oral suspension",
    "creme": "Cream",
    "salve": "Ointment",
    "gel": "Gel",
    "depotplaster": "Transdermal patch",
}
POLISH_COUNTRIES = {
    "polska": "Poland", "niemcy": "Germany", "hiszpania": "Spain", "włochy": "Italy",
    "francja": "France", "holandia": "Netherlands", "irlandia": "Ireland", "belgia": "Belgium",
    "austria": "Austria", "słowenia": "Slovenia", "czechy": "Czech Republic", "grecja": "Greece",
    "szwecja": "Sweden", "węgry": "Hungary", "dania": "Denmark", "malta": "Malta",
    "portugalia": "Portugal", "finlandia": "Finland", "bułgaria": "Bulgaria", "rumunia": "Romania",
    "łotwa": "Latvia", "litwa": "Lithuania", "estonia": "Estonia", "słowacja": "Slovakia",
    "chorwacja": "Croatia", "cypr": "Cyprus", "luksemburg": "Luxembourg", "islandia": "Iceland",
    "norwegia": "Norway", "szwajcaria": "Switzerland", "wielka brytania": "United Kingdom",
    "stany zjednoczone": "United States", "stany zjednoczone ameryki": "United States",
    "indie": "India", "chiny": "China", "japonia": "Japan", "izrael": "Israel", "kanada": "Canada",
    "turcja": "Turkey", "korea południowa": "South Korea", "republika korei": "South Korea",
    "australia": "Australia", "serbia": "Serbia", "ukraina": "Ukraine", "tajwan": "Taiwan",
}


def _translated(value: object, table: dict[str, str]) -> str:
    text = _clean(value)
    return table.get(text.lower(), text)


def _translated_list(value: object, table: dict[str, str]) -> str:
    """Several values on separate lines, each translated, repeats dropped."""
    return _unique(_translated(part, table) for part in str(value or "").splitlines())


# ------------------------------------------------------------ Netherlands

NETHERLANDS_URL = "https://www.geneesmiddeleninformatiebank.nl/metadata.csv"
NETHERLANDS_PAGE = "https://www.geneesmiddeleninformatiebank.nl/"


def read_netherlands(register: OpenRegister, path: Path) -> Iterator[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        yield from csv.DictReader(handle, delimiter="|")


def _netherlands_scope(procedure: str, number: str) -> str:
    if number.startswith("EU/"):
        return "Centralised"
    if re.match(r"^[A-Z]{2}/H/\d+", procedure):
        return "Mutual recognition / decentralised"
    return "National" if number else ""


def build_netherlands_rows(records: Iterable[dict[str, str]], fetched_at: str) -> Iterator[tuple[str, dict[str, Any]]]:
    for record in records:
        actives = [_clean(part) for part in str(record.get("WERKZAMESTOFFEN") or "").split("#") if _clean(part)]
        product = _clean(record.get("PRODUCTNAAM"))
        if not (actives or product):
            continue
        active = "; ".join(actives)
        kind = _clean(record.get("SOORT"))
        number = _clean(record.get("REGISTRATIENUMMER"))
        registration = f"{kind} {number}" if kind and not number.startswith("EU/") else number
        atc = _clean(record.get("ATC")).split(" - ", 1)[0]
        yield _indexed(national_match_text(active or product), {
            "substance": active,
            "active_substance": active,
            "source_substance": active,
            "product": product,
            "company": _clean(record.get("HANDELSVERGUNNINGHOUDER")),
            "country": "Netherlands",
            "region": "EU",
            "status": "Authorised",
            "authorisation_scope": _netherlands_scope(_clean(record.get("PROCEDURENUMMER")), number),
            "strength": _clean(record.get("POTENTIE")) or extract_strength(product),
            "dosage_form": _translated(record.get("FARMACEUTISCHEVORM"), DUTCH_FORMS),
            "route": _clean(record.get("TOEDIENINGSWEG")).replace("$SEMICOLON$", ";"),
            "atc_code": atc,
            "classification": _clean(record.get("AFLEVERSTATUS")).replace("$SEMICOLON$", ";"),
            "registration_number": registration,
            "registration_date": _date(record.get("INSCHRIJVINGSDATUM")),
            "source": "CBG Netherlands",
            "source_url": NETHERLANDS_PAGE,
            "product_url": _clean(record.get("PRODUCTDETAIL_LINK")),
            "smpc_url": _clean(record.get("SMPC_FILENAAM")),
            "pil_url": _clean(record.get("BIJSLUITER_FILENAAM")),
            "assessment_report_url": _clean(record.get("PAR_FILENAAM")),
            "document_type": "CBG Medicines Information Bank record",
            "last_checked": fetched_at,
        })


CBG_NETHERLANDS = add_register(OpenRegister(
    source="CBG Netherlands",
    country="Netherlands",
    region="EU",
    url=NETHERLANDS_URL,
    slug="cbg_netherlands",
    max_age_seconds=7 * 24 * 3600,
    build_rows=build_netherlands_rows,
    read_records=read_netherlands,
))


def run_cbg_netherlands_search(substance: str) -> list[dict[str, Any]]:
    return search_register(CBG_NETHERLANDS, substance)


# ----------------------------------------------------------------- Poland

POLAND_URL = "https://rejestry.ezdrowie.gov.pl/api/rpl/medicinal-products/public-pl-report/get-csv"
POLAND_PAGE = "https://rejestry.ezdrowie.gov.pl/rpl/search/public"
POLAND_PROCEDURES = {
    "NAR": "National",
    "DCP": "Decentralised",
    "MRP": "Mutual recognition",
    "CEN": "Centralised",
    "IR": "Parallel import",
}


def read_poland(register: OpenRegister, path: Path) -> Iterator[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        yield from csv.DictReader(handle, delimiter=";")


def _poland_packs(value: object) -> str:
    """ "05909991023652 ¦ Rpz ¦ 2\\n1 fiol. 5 ml" per pack -> "1 fiol. 5 ml; ..." """
    packs = []
    for block in re.split(r"\n(?=\d{8,14} ¦)", str(value or "")):
        lines = [line for line in block.splitlines() if "¦" not in line]
        if lines:
            packs.append(" ".join(lines))
    return _unique(packs)


def build_poland_rows(records: Iterable[dict[str, str]], fetched_at: str) -> Iterator[tuple[str, dict[str, Any]]]:
    for record in records:
        if _clean(record.get("Rodzaj preparatu")) != "Ludzki":  # veterinary products are not ours
            continue
        actives = [_clean(part) for part in str(record.get("Nazwa powszechnie stosowana") or "").split("+") if _clean(part)]
        product = _clean(record.get("Nazwa Produktu Leczniczego"))
        if not (actives or product):
            continue
        active = "; ".join(actives)
        validity = _clean(record.get("Ważność pozwolenia"))
        maker = _unique(str(record.get("Nazwa wytwórcy") or "").splitlines())
        procedure = _clean(record.get("Typ procedury"))
        strength = _clean(record.get("Moc"))
        form = _translated(record.get("Postać farmaceutyczna"), POLISH_FORMS)
        number = _clean(record.get("Numer pozwolenia"))
        if not number:
            # A centrally authorised product has no Polish permit number, and
            # one name covers every strength ("Zometa" 4 mg and 4 mg/5 ml), so
            # the strength and form are what tell its rows apart.
            product = " ".join(part for part in (product, strength, form) if part)
        yield _indexed(national_match_text(active or product, latin=True), {
            "substance": active,
            "active_substance": active,
            "source_substance": active,
            "product": product,
            "company": _clean(record.get("Podmiot odpowiedzialny")),
            "country": "Poland",
            "region": "EU",
            "status": "Authorised",
            "authorisation_scope": POLAND_PROCEDURES.get(procedure, procedure),
            "strength": strength,
            "dosage_form": form,
            "route": _clean(record.get("Droga podania - Gatunek - Tkanka - Okres karencji")),
            "pack_size": _poland_packs(record.get("Opakowanie")),
            "atc_code": _clean(record.get("Kod ATC")),
            "registration_number": number,
            "expiry_date": "Unlimited" if validity == "Bezterminowe" else _date(validity),
            "manufacturer_name": maker,
            "manufacturer_country": _translated_list(record.get("Kraj wytwórcy"), POLISH_COUNTRIES),
            "manufacturer_source": "Polish Register of Medicinal Products" if maker else "",
            "source": "URPL Poland",
            "source_url": POLAND_PAGE,
            "product_url": "",
            "smpc_url": _clean(record.get("Charakterystyka")),
            "pil_url": _clean(record.get("Ulotka")),
            "document_type": "Register of Medicinal Products record",
            "last_checked": fetched_at,
        })


URPL_POLAND = add_register(OpenRegister(
    source="URPL Poland",
    country="Poland",
    region="EU",
    url=POLAND_URL,
    slug="urpl_poland",
    max_age_seconds=24 * 3600,
    build_rows=build_poland_rows,
    read_records=read_poland,
))


def run_urpl_poland_search(substance: str) -> list[dict[str, Any]]:
    return search_register(URPL_POLAND, substance)


# ---------------------------------------------------------------- Czechia

CZECH_CATALOGUE = "https://opendata.sukl.cz/?q=katalog%2Fdatabaze-lecivych-pripravku-dlp"
CZECH_PAGE = "https://prehledy.sukl.cz/prehled_leciv.html"
# Registration states worth showing; the rest are treatment programmes,
# emergency permits and foods for special medical purposes, which are not
# marketing authorisations.
CZECH_STATUS = {
    "R": "Registered",
    "B": "Registered (previous presentation, may be sold for 6 months after a change)",
    "M": "Suspended",
    "K": "Centralised authorisation suspended",
    "C": "Cancelled, being withdrawn",
}
CZECH_PROCEDURES = {
    "NAR": "National",
    "CMS": "Mutual recognition / decentralised",
    "RMS": "Mutual recognition / decentralised",
    "EUR": "Centralised",
    "ORP": "Centralised",
    "SOU": "Parallel import",
    "SDI": "Parallel distribution",
}


def fetch_czech(register: OpenRegister, workdir: Path) -> Path:
    """The catalogue page names the current monthly ZIP; its name changes each month."""
    import requests

    from sources.open_registers import DOWNLOAD_TIMEOUT

    page = requests.get(CZECH_CATALOGUE, timeout=DOWNLOAD_TIMEOUT)
    page.raise_for_status()
    links = re.findall(r'href="(https://opendata\.sukl\.cz/soubory/[^"]+/DLP\d{8}\.zip)"', page.text)
    if not links:
        raise RuntimeError("SUKL's catalogue page names no DLP file")
    return download_file(register, sorted(links)[-1], workdir / "dlp.zip")


def _czech_table(archive: zipfile.ZipFile, name: str) -> Iterator[dict[str, str]]:
    with archive.open(name) as raw:
        yield from csv.DictReader(io.TextIOWrapper(raw, encoding="cp1250", newline=""), delimiter=";")


def read_czech(register: OpenRegister, path: Path) -> Iterator[dict[str, Any]]:
    """One record per registration and strength, its packs gathered together."""
    with zipfile.ZipFile(path) as archive:
        names = {row["KOD_LATKY"]: row.get("NAZEV_EN") or row.get("NAZEV_INN") or row.get("NAZEV") for row in _czech_table(archive, "dlp_latky.csv")}
        composition: dict[str, list[tuple[int, str]]] = defaultdict(list)
        for row in _czech_table(archive, "dlp_slozeni.csv"):
            if row.get("S") == "L" and names.get(row["KOD_LATKY"]):
                composition[row["KOD_SUKL"]].append((int(row.get("SQ") or 0), names[row["KOD_LATKY"]]))
        forms = {row["FORMA"]: row.get("NAZEV_EN") or row.get("NAZEV") for row in _czech_table(archive, "dlp_formy.csv")}
        routes = {row["CESTA"]: row.get("NAZEV_EN") or row.get("NAZEV") for row in _czech_table(archive, "dlp_cesty.csv")}
        countries = {row["ZEM"]: (row.get("NAZEV_EN") or row.get("NAZEV") or "").title() for row in _czech_table(archive, "dlp_zeme.csv")}
        holders = {row["ZKR_ORG"]: row.get("NAZEV") for row in _czech_table(archive, "dlp_organizace.csv")}

        grouped: dict[tuple[str, ...], dict[str, Any]] = {}
        for row in _czech_table(archive, "dlp_lecivepripravky.csv"):
            if row.get("REG") not in CZECH_STATUS:
                continue
            number = _clean(row.get("RC"))
            # A centrally authorised product's number names the pack in its
            # last part (EU/1/14/944/012); its packs are one product.
            if number.startswith("EU/"):
                number = "/".join(number.split("/")[:4])
            key = (number or row["KOD_SUKL"], row.get("NAZEV") or "", row.get("SILA") or "", row.get("FORMA") or "")
            record = grouped.get(key)
            if record is None:
                record = grouped[key] = {
                    **row,
                    "RC": number,
                    "actives": [name for _, name in sorted(composition.get(row["KOD_SUKL"], []))],
                    "form_en": forms.get(row.get("FORMA") or "", row.get("FORMA") or ""),
                    "route_en": routes.get(row.get("CESTA") or "", row.get("CESTA") or ""),
                    "holder": holders.get(row.get("DRZ") or "", row.get("DRZ") or ""),
                    "holder_country": countries.get(row.get("ZEMDRZ") or "", ""),
                    "packs": [],
                    "marketed": False,
                }
            record["packs"].append(_clean(row.get("DOPLNEK")))
            record["marketed"] = record["marketed"] or row.get("DODAVKY") == "1"
        yield from grouped.values()


def build_czech_rows(records: Iterable[dict[str, Any]], fetched_at: str) -> Iterator[tuple[str, dict[str, Any]]]:
    for record in records:
        active = "; ".join(dict.fromkeys(name.capitalize() for name in record["actives"]))
        name = _clean(record.get("NAZEV"))
        strength = _clean(record.get("SILA"))
        product = f"{name} {strength}".strip()
        if not (active or product):
            continue
        status = CZECH_STATUS[record["REG"]]
        if record["REG"] == "R" and record["marketed"]:
            status = "Registered, supplied in the last six months"
        validity = _clean(record.get("V_PLATDO"))
        yield _indexed(national_match_text(active or name), {
            "substance": active,
            "active_substance": active,
            "source_substance": active,
            "product": product,
            "company": _clean(record.get("holder")),
            "country": "Czech Republic",
            "region": "EU",
            "status": status,
            "authorisation_scope": CZECH_PROCEDURES.get(record.get("REG_PROC") or "", record.get("REG_PROC") or ""),
            "strength": strength,
            "dosage_form": _clean(record.get("form_en")).capitalize(),
            "route": _clean(record.get("route_en")),
            "pack_size": _unique(record["packs"]),
            "atc_code": _clean(record.get("ATC_WHO")),
            "registration_number": _clean(record.get("RC")),
            "expiry_date": "Unlimited" if record.get("NEOMEZ") == "X" else _date(validity),
            "source": "SUKL Czech Republic",
            "source_url": CZECH_PAGE,
            "product_url": "",
            "document_type": "SUKL medicinal products database (DLP) record",
            "last_checked": fetched_at,
        })


SUKL_CZECH = add_register(OpenRegister(
    source="SUKL Czech Republic",
    country="Czech Republic",
    region="EU",
    url=CZECH_CATALOGUE,
    slug="sukl_czech",
    max_age_seconds=7 * 24 * 3600,
    build_rows=build_czech_rows,
    fetch=fetch_czech,
    read_records=read_czech,
))


def run_sukl_czech_search(substance: str) -> list[dict[str, Any]]:
    return search_register(SUKL_CZECH, substance)


# ---------------------------------------------------------------- Denmark

DENMARK_URL = "https://laegemiddelstyrelsen.dk/ftp-upload/ListeOverGodkendteLaegemidler.xlsx"
DENMARK_PAGE = (
    "https://laegemiddelstyrelsen.dk/en/licensing/licensing-of-medicines/"
    "lists-of-authorised-and-deregistered-medicines/"
)
DENMARK_PROCEDURES = {
    "National": "National",
    "MRP": "Mutual recognition",
    "Decentral procedure": "Decentralised",
    "Central": "Centralised",
    "Par-Imp": "Parallel import",
}
DENMARK_COLUMNS = (
    "Drugid", "Navn", "Lægemiddelform", "Styrketekst", "AktiveSubstanser", "MftIndehaver",
    "ATC-kode", "Godkendt procedure", "Godkendt rolle", "Registreringsdato", "Er i Medicinpriser",
)


def fetch_denmark(register: OpenRegister, workdir: Path) -> Path:
    # openpyxl reads a workbook by its extension.
    return download_file(register, register.url, workdir / "dkma_authorised.xlsx")


def read_denmark(register: OpenRegister, path: Path) -> Iterator[dict[str, Any]]:
    """The first sheet; its header row names the columns, which sit among blank ones."""
    import openpyxl

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows)
        positions = {name: header.index(name) for name in DENMARK_COLUMNS if name in header}
        missing = [name for name in DENMARK_COLUMNS if name not in positions]
        if missing:
            raise RuntimeError(f"DKMA's list has changed: no {', '.join(missing)} column")
        for row in rows:
            if row and row[positions["Drugid"]]:
                yield {name: row[index] for name, index in positions.items()}
    finally:
        workbook.close()


def build_denmark_rows(records: Iterable[dict[str, Any]], fetched_at: str) -> Iterator[tuple[str, dict[str, Any]]]:
    for record in records:
        if str(record.get("ATC-kode") or "").startswith("Q"):  # ATCvet: a veterinary medicine
            continue
        active = _unique(str(record.get("AktiveSubstanser") or "").split(","))
        name = _clean(record.get("Navn"))
        if not (active or name):
            continue
        strength = _clean(record.get("Styrketekst"))
        form_da = _clean(record.get("Lægemiddelform"))
        # The list has no authorisation number and one name serves every
        # strength, so the name, strength and form together tell rows apart.
        product = " ".join(part for part in (name, strength, form_da) if part)
        procedure = _clean(record.get("Godkendt procedure"))
        holder = _clean(record.get("MftIndehaver"))
        if procedure == "Par-Imp":
            # Each parallel importer is listed under the original's name.
            product = f"{product} (parallel import: {holder})"
        yield _indexed(national_match_text(active or name), {
            "substance": active,
            "active_substance": active,
            "source_substance": active,
            "product": product,
            "company": holder,
            "country": "Denmark",
            "region": "EU",
            "status": "Authorised, on the price list" if _clean(record.get("Er i Medicinpriser")) == "Ja" else "Authorised",
            "authorisation_scope": DENMARK_PROCEDURES.get(procedure, procedure),
            "strength": strength,
            "dosage_form": _translated(form_da, DANISH_FORMS),
            "atc_code": _clean(record.get("ATC-kode")),
            "registration_number": "",
            "registration_date": _date(record.get("Registreringsdato")),
            "source": "DKMA Denmark",
            "source_url": DENMARK_PAGE,
            "product_url": "",
            "document_type": "DKMA list of authorised medicines",
            "last_checked": fetched_at,
        })


DKMA_DENMARK = add_register(OpenRegister(
    source="DKMA Denmark",
    country="Denmark",
    region="EU",
    url=DENMARK_URL,
    slug="dkma_denmark",
    max_age_seconds=24 * 3600,
    build_rows=build_denmark_rows,
    fetch=fetch_denmark,
    read_records=read_denmark,
))


def run_dkma_denmark_search(substance: str) -> list[dict[str, Any]]:
    return search_register(DKMA_DENMARK, substance)
